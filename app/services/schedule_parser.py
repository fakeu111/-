"""把教务系统导出的 Word(.docx) 课表解析为结构化上课安排。

已确认的格式（每个课程块占两行，多门课依次换行）：
    课程名 班号
    (a~b周) (p-q节) 教室 老师
- 一张“周一~周日 × 第 1~12 节”网格，分上午/下午/晚；
- 同一格可含多门课；同一门课不同周可换老师/教室（分别成条）。
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import docx

WEEK_RE = re.compile(r"(\d+)\s*[~\-–]\s*(\d+)\s*周")
PERIOD_RE = re.compile(r"(\d+)\s*[\-–~]\s*(\d+)\s*节")
NAME_RE = re.compile(r"^(.+?)\s+(\d{3,4})$")
WEEK_PAREN_RE = re.compile(r"\([^)]*周[^)]*\)")
PERIOD_PAREN_RE = re.compile(r"\([^)]*节[^)]*\)")
REMARK_RE = re.compile(r"\d\s*[.．]\s*\S")


def _parse_course(name_chunk: str, fields_chunk: str, day: int) -> dict[str, Any]:
    m = NAME_RE.match(name_chunk)
    cur: dict[str, Any] = {
        "name": m.group(1) if m else name_chunk,
        "code": m.group(2) if m else "",
        "day": day,
    }
    wm = WEEK_RE.search(fields_chunk)
    if wm:
        cur["ws"] = int(wm.group(1))
        cur["we"] = int(wm.group(2))
    pm = PERIOD_RE.search(fields_chunk)
    if pm:
        cur["ps"] = int(pm.group(1))
        cur["pe"] = int(pm.group(2))

    rest = WEEK_PAREN_RE.sub("", fields_chunk)
    rest = PERIOD_PAREN_RE.sub("", rest)
    tokens = rest.split()
    if tokens:
        cur["room"] = tokens[0]
        cur["teacher"] = tokens[-1]
    return cur


def _parse_block(raw: str, day: int) -> list[dict[str, Any]]:
    parts = [
        p.strip()
        for p in raw.replace("\xa0", " ").split("\n")
        if p.strip()
    ]
    courses: list[dict[str, Any]] = []
    i = 0
    while i < len(parts):
        if NAME_RE.match(parts[i]):
            fields = parts[i + 1] if i + 1 < len(parts) else ""
            courses.append(_parse_course(parts[i], fields, day))
            i += 2
        else:
            i += 1
    return courses


def parse_docx(file: str | Path) -> tuple[list[dict[str, Any]], int]:
    """解析课表文件，返回 (上课安排列表, 备注课程数)。"""
    doc = docx.Document(str(file))

    # 找到表头含“星期一”的那张表
    table = None
    for tbl in doc.tables:
        joined = " ".join(c.text for c in tbl.rows[0].cells)
        if "星期一" in joined:
            table = tbl
            break
    if table is None:
        raise ValueError("未在文档中找到课程网格表（表头需包含“星期一”）。")

    header = [c.text.replace("\xa0", "").strip() for c in table.rows[0].cells]
    day_col: dict[int, int] = {}
    for col, text in enumerate(header):
        for d, name in enumerate(
            ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"], start=1
        ):
            if text == name:
                day_col[d] = col

    # 识别节次列：前几列中能取到 1..12 数字的那一列
    sample = table.rows[1].cells
    period_col = None
    for col in range(min(3, len(sample))):
        if sample[col].text.strip().isdigit():
            period_col = col
            break
    if period_col is None:
        raise ValueError("未能识别节次列。")

    courses: list[dict[str, Any]] = []
    last_raw: dict[int, str] = {}
    for row in table.rows[1:]:
        cells = row.cells
        period_text = cells[period_col].text.strip()
        if not period_text.isdigit():
            continue
        for day, col in day_col.items():
            raw = cells[col].text.replace("\xa0", " ").strip()
            if not raw or raw == last_raw.get(day):
                continue
            last_raw[day] = raw
            courses.extend(_parse_block(cells[col].text, day))

    # 备注课程数：顶部段落“备注:”后的编号条目
    remark_count = 0
    for p in doc.paragraphs:
        if "备注" in p.text:
            remark_count += len(REMARK_RE.findall(p.text))

    return courses, remark_count
